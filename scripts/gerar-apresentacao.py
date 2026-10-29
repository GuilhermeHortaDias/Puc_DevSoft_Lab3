"""Gera PPTX editável e PDF. Dependências: python-pptx e reportlab<5 (Python 3.9)."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
import textwrap
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
FONT_DIR = Path("/System/Library/Fonts/Supplemental")
PDF_FONT, PDF_BOLD = "Helvetica", "Helvetica-Bold"
if (FONT_DIR / "Arial.ttf").exists():
    pdfmetrics.registerFont(TTFont("Arial", str(FONT_DIR / "Arial.ttf")))
    pdfmetrics.registerFont(TTFont("Arial-Bold", str(FONT_DIR / "Arial Bold.ttf")))
    PDF_FONT, PDF_BOLD = "Arial", "Arial-Bold"

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/apresentacao'
OUT.mkdir(exist_ok=True)
prs = Presentation()
prs.slide_width, prs.slide_height = Pt(960), Pt(540)
pdf = canvas.Canvas(str(OUT / 'Sistema-de-Moeda-Estudantil.pdf'), pagesize=(960,540))
pdf.setTitle('Sistema de Moeda Estudantil — Laboratório 3')
BG, WHITE, MUTED, ACCENT, PANEL = '101D2C','F4F7FA','B6C7D9','56D9B0','1B2D42'
slide = None

def box(x,y,w,h,color):
    sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(color); sh.line.fill.background()
    pdf.setFillColor(HexColor('#'+color)); pdf.rect(x,540-y-h,w,h,fill=1,stroke=0)

def txt(x,y,w,text,size=20,color=WHITE,bold=False):
    # Explicit wrapping keeps both exports aligned and every element editable.
    lines=[]
    for line in text.split('\n'):
        lines.extend(textwrap.wrap(line,width=max(1,int(w/(size*.53))),break_long_words=False) or [''])
    tf=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(len(lines)*size*1.27+8)).text_frame
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Arial'; p.font.size=Pt(size); p.font.bold=bold
        p.font.color.rgb=RGBColor.from_string(color); p.space_after=Pt(0); p.line_spacing=1.15
    pdf.setFillColor(HexColor('#'+color)); pdf.setFont(PDF_BOLD if bold else PDF_FONT,size)
    for i,line in enumerate(lines): pdf.drawString(x,540-y-size-i*size*1.15,line)

def start(n,tag,title,sub=''):
    global slide
    n = len(prs.slides) + 1
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    box(0,0,960,540,BG); box(40,38,35,4,ACCENT)
    txt(86,28,810,tag.upper(),12,ACCENT,True)
    txt(40,72,880,title,34,WHITE,True)
    if sub: txt(40,130,880,sub,18,MUTED)
    box(40,496,880,1,PANEL)
    txt(40,510,820,'PUC MINAS  /  PROJETO DE SOFTWARE  /  LABORATÓRIO 3',10,MUTED)
    txt(885,507,40,f'{n:02d}',14,ACCENT,True)

def end(notes):
    slide.notes_slide.notes_text_frame.text=notes
    pdf.showPage()

def cards(items,y=192,h=240):
    gap=16; w=(880-gap*(len(items)-1))/len(items)
    for i,(head,body) in enumerate(items):
        x=40+i*(w+gap); box(x,y,w,h,PANEL)
        txt(x+20,y+18,w-40,head,22,ACCENT,True)
        txt(x+20,y+70,w-40,body,19)

start(1,'Apresentação do projeto','Sistema de Moeda Estudantil','Reconhecimento acadêmico que se transforma em benefícios.')
txt(40,218,600,'Mérito → moedas → vantagens',32,ACCENT,True)
txt(40,302,850,'Guilherme Horta Dias\nLucas Batista Duarte\nRafael Abras',22)
txt(40,424,850,'Engenharia de Software · Profa. Milena Menezes Adão · 2026/2',17,MUTED)
end('Tempo: 30 s. Apresentar o grupo e o objetivo: reconhecer alunos por meio de moedas virtuais, com benefícios oferecidos por empresas parceiras.')

start(2,'Contexto','Do reconhecimento ao benefício','Três perfis participam de um único ciclo de mérito estudantil.')
cards([('Professor','Reconhece o aluno e registra o motivo do envio de moedas.'),('Aluno','Acompanha saldo e extrato e troca moedas por vantagens.'),('Empresa parceira','Publica ofertas e recebe o cupom do resgate por email.')])
end('Tempo: 40 s. Explicar o papel de cada perfil. As moedas são virtuais, usadas no contexto do laboratório. O professor não se cadastra publicamente: seus dados vêm da instituição.')

start(4,'Regras de negócio','Como as moedas circulam','Exemplo ilustrativo de reconhecimento e resgate.')
cards([('1.000 moedas','O professor recebe crédito a cada semestre, somado ao saldo restante.'),('Reconhecimento','O professor envia 100 moedas e informa o motivo. O aluno recebe o crédito.'),('Resgate','Uma vantagem de 60 moedas gera um cupom. O aluno fica com 40 moedas.')],h=250)
end('Tempo: 50 s. Os valores 100 e 60 são exemplos, não dados reais. O aluno começa com zero; a empresa não possui carteira. O saldo não pode ficar negativo. O professor consulta alunos ativos da mesma instituição, decisão D09 do grupo.')

start(5,'Arquitetura MVC','Responsabilidades bem definidas','A interface apresenta; os serviços aplicam as regras de negócio.')
for i,(h,b) in enumerate([('React + TypeScript','Telas e chamadas HTTP'),('Controllers','HTTP, DTOs e validação'),('Services','Regras e transações'),('JPA + PostgreSQL','Entidades e persistência')]):
    x=40+i*224; box(x,210,208,150,PANEL); txt(x+14,230,180,h,21,ACCENT,True); txt(x+14,292,180,b,17)
    if i<3: txt(x+208,267,20,'›',19,ACCENT,True)
txt(40,400,870,'Spring Boot no backend · Flyway versiona o esquema do banco.',21,MUTED)
end('Tempo: 1 min. Explicar MVC e a separação em camadas. Controllers recebem DTOs; services verificam autorização e regras dentro de transações; repositories acessam PostgreSQL por Spring Data JPA. Hibernate valida o esquema criado pelo Flyway.')

start(6,'Modelo de dados','Identidade, saldo e histórico','Visão resumida das relações; o ER completo está na documentação.')
cards([('Identidade','Usuário vincula aluno, professor ou empresa.\nAluno e professor têm instituição.'),('Moedas','Aluno e professor têm conta.\nMovimentações registram créditos, envios e resgates.'),('Benefícios','Empresa publica vantagem.\nResgate guarda cupom.\nNotificação registra o email.')],h=264)
end('Tempo: 50 s. Referência: docs/modelagem/persistencia.md e diagramas/modelo-er.png. Aluno, professor e empresa associam o UUID de Usuario como PK e FK, sem herança JPA. O cupom fica na movimentação de resgate. A descrição e o custo resgatados são preservados no histórico.')

start(7,'Integridade','Uma operação, um registro consistente','Saldo, movimentação e intenção de email são confirmados juntos.')
cards([('Saldo protegido','Bloqueio de contas e validação impedem débito acima do saldo.'),('Sem duplicidade','Uma identificação por comando evita repetir o débito ao reenviar a operação.'),('Histórico preservado','Inativação conserva cadastros e movimentações já realizadas.')])
end('Tempo: 45 s. A unicidade de solicitante e requisicaoId protege o comando. Repetir o mesmo comando retorna o registro original; reutilizar o identificador com outro comando gera conflito. Crédito semestral também tem proteção contra duplicidade. Falhas causam rollback da operação.')

start(8,'Acesso e segurança','Cada perfil acessa suas funções','Autorização no backend acompanha a navegação da interface.')
cards([('Autenticação','Senhas com BCrypt.\nSessão com cookie HttpOnly e proteção CSRF.'),('Autorização','Aluno e empresa mantêm apenas seu próprio cadastro. Professor reconhece alunos.'),('Validação','DTOs e Bean Validation.\nEntidades e hashes de senha não são expostos pela API.')])
end('Tempo: 40 s. Explicar que esconder botões não basta: os serviços validam o titular e o perfil. Cadastro e login são públicos; demais operações exigem sessão. Professores são importados por arquivo institucional privado.')

start(9,'Emails e cupons','O resgate chega às duas partes','Aluno e empresa recebem o mesmo código de cupom.')
for i,(h,b) in enumerate([('Resgate confirmado','Banco registra débito, cupom e notificações.'),('Worker de email','Processa o envio e tenta novamente após falha.'),('Mailpit local','Captura mensagens em localhost:8025.')]):
    y=180+i*95; box(40,y,880,81,PANEL); txt(60,y+13,260,h,21,ACCENT,True); txt(350,y+14,540,b,20)
end('Tempo: 45 s. O envio SMTP ocorre após a confirmação da operação. Se falhar, apenas a notificação é repetida, sem novo débito. Mailpit captura emails locais; um provedor SMTP no .env permite entrega externa. A conferência do cupom é presencial, sem baixa digital nesta versão.')

start(0,'CRUDs de aluno e empresa','Criar, consultar, editar e inativar','A mesma estrutura de operações atende aos dois perfis.')
cards([('Cadastro e consulta','POST cria o cadastro.\nGET consulta os dados do titular autenticado.'),('Atualização','PUT altera os dados.\nA empresa mantém nome e email; o aluno mantém seus dados acadêmicos.'),('Inativação','DELETE inativa o cadastro e encerra a sessão.\nO histórico é conservado.')],h=250)
end('Tempo: 1 min. Mostrar os endpoints em CadastrosController.java. Alunos usam /api/alunos e empresas /api/empresas. IDs aparecem nas operações individuais. Na demonstração, criar, consultar e editar nos dois perfis. Usar cadastros descartáveis para demonstrar inativação sem interromper o fluxo principal.')

start(0,'Modelagem e código','O modelo orienta a implementação','Correspondências que podem ser conferidas no repositório.')
cards([('Classes do domínio','Usuario, Aluno e Conta aparecem no modelo e nas entidades do pacote cadastro.'),('Componentes MVC','React → controllers → services → repositories. A separação aparece no código.'),('Modelo ER','Entidades e relações acompanham as migrações SQL em db/migration.')],h=250)
end('Tempo: 1 min 30 s. Abrir docs/modelagem/diagramas/classes.svg, componentes.svg e modelo-er.svg. Comparar Usuario, Aluno e Conta com as entidades do pacote cadastro. Mostrar o limite entre frontend e backend no diagrama de componentes. No ER, localizar conta e movimentacao, comparando com V3__moedas_vantagens_notificacoes.sql. As imagens e fontes PlantUML estão versionadas. O PDF da Sprint 1 é histórico; os modelos em docs/modelagem são os atuais.')

start(0,'Tutorial / 1 de 4','React: da tela à requisição','Uma alteração de cadastro percorre componentes e cliente HTTP.')
cards([('1. Localizar a tela','Abrir CadastroForm.tsx.\nIdentificar campos, estado do formulário e ação de envio.'),('2. Seguir o envio','Abrir api.ts.\nObservar os tipos de dados, a chamada HTTP e o tratamento de erro.'),('3. Observar no navegador','Editar um cadastro.\nEm Rede / Network, conferir a requisição, o status e a resposta.')],h=260)
end('Tempo: 1 min 30 s. Tutorial ao vivo: abrir frontend/src/CadastroForm.tsx e frontend/src/api.ts no editor. Mostrar a relação entre campo e dados enviados. No navegador autenticado, abrir as ferramentas de desenvolvedor e editar um campo. Inspecionar o PUT sem expor cookies ou credenciais. Explicar que TypeScript descreve os dados e React atualiza a interface a partir do estado.')

start(0,'Tutorial / 2 de 4','Spring Boot: da rota à regra','Seguir uma requisição pelo controller e pelo serviço.')
cards([('Rota HTTP','Em CadastrosController, localizar @PostMapping("/alunos") e @RequestBody.'),('Validação e resposta','@Valid valida a entrada.\nResponseEntity.created devolve 201 e o endereço do novo recurso.'),('Regra de negócio','Em CadastrosService, localizar cadastrarAluno e acompanhar a transação.')],h=260)
end('Tempo: 1 min 30 s. Abrir backend/src/main/java/br/pucminas/moedaestudantil/cadastro/CadastrosController.java e CadastrosService.java. Mostrar o método cadastrarAluno. Explicar injeção de dependência pelo construtor, DTO de entrada, validação e a delegação para service. A lógica fica no serviço para preservar as responsabilidades MVC.')

start(0,'Tutorial / 3 de 4','JPA e Flyway: do objeto à tabela','O ORM acessa dados; as migrações definem o esquema.')
cards([('Entidade JPA','Abrir Aluno.java.\nLocalizar @Entity, identificador e relações com usuário e instituição.'),('Repository','Abrir AlunoRepository.\nJpaRepository fornece operações básicas sobre Aluno e UUID.'),('Migração SQL','Abrir V1__cadastros.sql.\nComparar tabelas, chaves e restrições com as entidades.')],h=260)
end('Tempo: 1 min 30 s. Abrir as três fontes indicadas. Mostrar que JpaRepository<Aluno, UUID> oferece persistência sem implementar manualmente cada consulta SQL. Mostrar uma FK na migração. Flyway aplica versões em sequência e Hibernate valida o mapeamento; alterações futuras usam novas migrações. PostgreSQL é o banco da aplicação atual.')

start(0,'Tutorial / 4 de 4','Docker e Mailpit: observar o email','Um ambiente local reproduzível para a demonstração.')
box(40,185,880,75,PANEL)
txt(60,207,840,'docker compose up -d db mailpit',25,ACCENT,True)
cards([('Serviços locais','compose.yaml define PostgreSQL e Mailpit.\nSMTP local: porta 1025.'),('Configuração','O backend usa SMTP_HOST e SMTP_PORT do ambiente.\nNão exibir senhas.'),('Conferência','Abrir localhost:8025.\nReconhecer um aluno e observar a notificação recebida.')],y=282,h=197)
end('Tempo: 1 min 30 s. Abrir compose.yaml e executar o comando exibido, com Docker já iniciado. Explicar container e mapeamento de portas. Mostrar apenas as variáveis SMTP sem credenciais. A aplicação roda fora dos containers nesta configuração. Mailpit recebe SMTP em 1025 e exibe mensagens em 8025. A entrega externa depende de um provedor SMTP; o teste local é feito no Mailpit.')

start(10,'Demonstração ao vivo','Percorrer o ciclo na aplicação','Aplicação: localhost:5173  ·  Emails: localhost:8025')
steps=[('01','Aluno e empresa','Criar, consultar e editar os dois perfis; publicar vantagem.'),('02','Professor','Entrar, escolher aluno e enviar moedas com um motivo.'),('03','Aluno','Conferir o extrato e resgatar a vantagem.'),('04','Mailpit','Abrir os emails do resgate e comparar o código do cupom.')]
for i,(num,h,b) in enumerate(steps):
    y=181+i*72; txt(40,y,55,num,26,ACCENT,True); txt(105,y,245,h,21,WHITE,True); txt(360,y,550,b,19,MUTED)
end('Tempo: 5 min. Demonstrar também inativação com cadastros descartáveis de aluno e empresa. Preparar previamente aluno, empresa e professor fictícios. Login do professor: professor.demo; consultar a senha no arquivo local, sem mostrá-la nos slides. A vantagem deve custar menos que o saldo disponível. Alternar as sessões ao demonstrar os perfis. Confirmar backend, frontend, PostgreSQL e Mailpit antes da apresentação. Os endereços são da máquina que executa a aplicação.')

start(12,'Encerramento','Do mérito ao benefício, com histórico','Cadastro → reconhecimento → resgate → cupom → email')
txt(40,220,865,'Três perfis conectados por um fluxo de moedas rastreável.',32,ACCENT,True)
txt(40,334,865,'MVC · PostgreSQL · transações · autorização · notificações',23)
txt(40,423,865,'Demonstração e perguntas',25,WHITE,True)
end('Tempo: 20 s. Retomar o objetivo do sistema e abrir para perguntas. Referências internas: README.md; docs/modelagem/persistencia.md; docs/modelagem/requisitos.md; código em backend/ e frontend/.')

prs.core_properties.title='Sistema de Moeda Estudantil'
prs.core_properties.author='Guilherme Horta Dias; Lucas Batista Duarte; Rafael Abras'
prs.core_properties.subject='PUC Minas — Projeto de Software — Laboratório 3'
prs.save(OUT / 'Sistema-de-Moeda-Estudantil.pptx')
pdf.save()
print(f'Gerados {len(prs.slides)} slides em {OUT}')
